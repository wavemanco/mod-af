#!/usr/bin/env python3
"""Build jargone's word lists into hooks/data.ts, and refresh the build record in data/NOTICE.md.

Run it from anywhere:

    python3 jargone/scripts/build_data.py

By default it downloads every source from the URLs below. To build from copies you already have,
give their paths:

    python3 jargone/scripts/build_data.py --familiarity cleaned.xlsx --gsa gsa.md --eighteen-f 18f.md \
        --govuk a-to-z.html --iso-lexicon lexicon.ts --iso-audit audit-corpus.ts

It needs Python 3.9 or later and openpyxl (to read the familiarity spreadsheet). Everything else is
the standard library.

What it writes:
- hooks/data.ts: three strings, FAMILIARITY, REPLACEMENTS and EXCEPTIONS (the plugin's loader and its
  test kit import code files only, so the lists ship as a generated code module).
- data/NOTICE.md: only the part between <!-- build:start --> and <!-- build:end -->: the build date,
  each source's address and fingerprint, the counts, and every entry this script changed or left out.
  The rest of NOTICE.md is written by hand.

The technical-exceptions list is jargone's own. It lives in EXCEPTIONS below, one reason per word, and
this is the only place to change it.
"""

from __future__ import annotations

import argparse
import datetime as _dt
import hashlib
import html.parser
import json
import re
import sys
import unicodedata
import urllib.request
from pathlib import Path

# ---------------------------------------------------------------------------------------------
# Sources. GitHub files are pinned to the commit they were built from, so a rebuild gets the same
# bytes. The GOV.UK page and the OSF file cannot be pinned; their fingerprints (SHA-256) are
# recorded in NOTICE.md instead, so a change upstream shows up as a changed fingerprint.
# ---------------------------------------------------------------------------------------------

SOURCES = {
    "familiarity": {
        "flag": "--familiarity",
        "name": "AI-generated estimates of English word familiarity (cleaned list)",
        "url": "https://osf.io/download/bafpv/",
        "binary": True,
    },
    "gsa": {
        "flag": "--gsa",
        "name": "Federal Plain Language Guidelines: Use simple words and phrases",
        "url": "https://raw.githubusercontent.com/GSA/plainlanguage.gov/"
        "fd7694740f19c0ed20ed71c2dd1dc92699e920fa/_pages/guidelines/words/use-simple-words-phrases.md",
    },
    "18f": {
        "flag": "--eighteen-f",
        "name": "18F Content Guide: Use plain language (Words to avoid)",
        "url": "https://raw.githubusercontent.com/18F/content-guide/"
        "1b1723d3d5b8f91d92c16487c88b56265dc0ec3a/_pages/our-approach/plain-language.md",
    },
    "govuk": {
        "flag": "--govuk",
        "name": "GOV.UK A to Z style guide: Words to avoid",
        "url": "https://guidance.publishing.service.gov.uk/writing-to-gov-uk-standards/style-guides/a-to-z-style-guide/",
    },
    "iso-lexicon": {
        "flag": "--iso-lexicon",
        "name": "iso-24495 plugin: lib/lexicon.ts (COMPLEX_WORDS)",
        "url": "https://raw.githubusercontent.com/GaZmagik/iso-24495/"
        "5951eb7d985018459c62428378926c07f4199794/skills/iso-24495-4/scripts/lib/lexicon.ts",
    },
    "iso-audit": {
        "flag": "--iso-audit",
        "name": "iso-24495 plugin: audit-corpus.ts (WORDY_PHRASES, DOUBLETS)",
        "url": "https://raw.githubusercontent.com/GaZmagik/iso-24495/"
        "5951eb7d985018459c62428378926c07f4199794/skills/iso-24495-4/scripts/audit-corpus.ts",
    },
}

# The short names that go in each replacement entry's `source` field, in the order sources merge.
SRC_GSA = "plainlanguage.gov"
SRC_18F = "18F"
SRC_GOVUK = "GOV.UK"
SRC_ISO = "iso-24495"


class BuildError(Exception):
    """A source no longer looks the way this script expects. Read the source, then fix the script."""


# ---------------------------------------------------------------------------------------------
# Small text helpers
# ---------------------------------------------------------------------------------------------

_QUOTES = str.maketrans({"’": "'", "‘": "'", "“": '"', "”": '"', " ": " "})


def clean(s: str) -> str:
    """Straight quotes, single spaces, no ends, NFC."""
    return unicodedata.normalize("NFC", re.sub(r"\s+", " ", s.translate(_QUOTES)).strip())


def split_top(s: str, sep: str) -> list[str]:
    """Split on sep, but not inside round brackets."""
    parts, depth, start, i = [], 0, 0, 0
    while i < len(s):
        c = s[i]
        if c == "(":
            depth += 1
        elif c == ")":
            depth = max(0, depth - 1)
        elif depth == 0 and s.startswith(sep, i):
            parts.append(s[start:i])
            start = i + len(sep)
            i = start
            continue
        i += 1
    parts.append(s[start:])
    return parts


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


# ---------------------------------------------------------------------------------------------
# Familiarity list
# ---------------------------------------------------------------------------------------------

# One or more letters, with apostrophes only inside (don't, o'clock). No digits, spaces, hyphens,
# dots, slashes or underscores (the source writes some phrases with underscores, like air_conditioner).
WORD = re.compile(r"[^\W\d_]+(?:'[^\W\d_]+)*")
# A form ending in 's (possessives such as abbey's, and it's, that's). A checker removes a final 's
# before looking a word up, so these entries add nothing and are left out.
ENDS_IN_APOSTROPHE_S = re.compile(r".+'s")


def build_familiarity(data: bytes) -> tuple[list[tuple[str, str]], dict]:
    import io

    try:
        import openpyxl  # type: ignore
    except ImportError as e:  # pragma: no cover - depends on the machine
        raise BuildError("openpyxl is needed to read the familiarity spreadsheet: pip install openpyxl") from e

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    if "Values" not in wb.sheetnames:
        raise BuildError(f"familiarity: no 'Values' sheet (found {wb.sheetnames})")
    rows = wb["Values"].iter_rows(values_only=True)
    header = [str(h) for h in next(rows)]
    for col in ("Word", "GPT_Fam_probs", "Type"):
        if col not in header:
            raise BuildError(f"familiarity: column {col!r} missing (found {header})")
    i_word, i_fam, i_type = header.index("Word"), header.index("GPT_Fam_probs"), header.index("Type")

    stats = {
        "rows": 0,
        "phrases": 0,  # Type MWE
        "hyphenated": 0,  # Type WhW
        "other_type": 0,
        "not_letters": 0,  # numbers, abbreviations with dots, slashes, underscores, symbols
        "ends_in_s": 0,
        "same_word_other_case": 0,
    }
    best: dict[str, tuple[bool, float]] = {}  # word -> (was lower case in the source, rating)
    for row in rows:
        if row is None or row[i_word] is None:
            continue
        stats["rows"] += 1
        kind = row[i_type]
        if kind == "MWE":
            stats["phrases"] += 1
            continue
        if kind == "WhW":
            stats["hyphenated"] += 1
            continue
        if kind != "W":
            stats["other_type"] += 1
            continue
        word = clean(str(row[i_word]))
        if not WORD.fullmatch(word):
            stats["not_letters"] += 1
            continue
        low = word.lower()
        if ENDS_IN_APOSTROPHE_S.fullmatch(low):  # on the lower-cased word, so ADAMS'S goes too
            stats["ends_in_s"] += 1
            continue
        rating = float(row[i_fam])
        was_lower = word == low
        if low in best:
            stats["same_word_other_case"] += 1
            old_lower, old_rating = best[low]
            # Prefer the rating of the entry that was already lower case; otherwise the higher one.
            if (was_lower, rating) <= (old_lower, old_rating):
                continue
        best[low] = (was_lower, rating)

    lines = []
    for word in sorted(best):
        text = f"{best[word][1]:.1f}"
        if not (1.0 <= float(text) <= 7.0):
            raise BuildError(f"familiarity: rating {text} for {word!r} is outside 1 to 7")
        lines.append((word, text))
    ratings = [float(t) for _, t in lines]
    stats["kept"] = len(lines)
    stats["lowest"] = f"{min(ratings):.1f}"
    stats["highest"] = f"{max(ratings):.1f}"
    stats["at_or_above_6"] = sum(1 for r in ratings if r >= 6.0)
    return lines, stats


# ---------------------------------------------------------------------------------------------
# Replacement list: one parser per source, then a merge.
#
# Each parser turns a source line into (froms, to, unless). `to` items in round brackets are advice,
# not words to paste in; "(omit)" means "leave it out". This is the convention the plainlanguage.gov
# table already uses.
#
# Where a line cannot be read by the general rule, or the general rule would give a wrong result, a
# FIX names the exact source line and what to do, with the reason. Every fix must match a source
# line, or the build stops: a source that changed upstream must be read again by a person.
# ---------------------------------------------------------------------------------------------


def fix(action: str, reason: str, froms=None, to=None, unless=None) -> dict:
    return {"action": action, "reason": reason, "froms": froms, "to": to, "unless": unless}


GSA_FIXES = {
    "a and/or b": fix("set", "the source writes the pattern with letters; the phrase to find is 'and/or'",
                      froms=["and/or"], to=["a or b or both"]),
    "for example,______etc.": fix("drop", "a pattern ('for example' used with 'etc.'), not a word or phrase to find"),
    "said, some, such": fix("drop", "these are everyday words; the source means them only as legal pointing "
                                    "words ('said contract'), which a word list cannot tell apart from everyday use, "
                                    "and the same table offers 'some' as a replacement for 'a number of'"),
    "not later than 10 May": fix("set", "the source gives two dated examples; the phrase is 'not later than'. "
                                         "Only 'by' is kept: the source's 'before 11 May' moves the date by a day, "
                                         "and 'before' alone with the same date would change the deadline",
                                 froms=["not later than"], to=["by"]),
    "not later than 1600": fix("set", "merged with the other 'not later than' example",
                               froms=["not later than"], to=["by"]),
    "/ (slash)": fix("drop", "a punctuation mark, not a word or phrase"),
    "is responsible for": fix("set", "the source writes '(omit) handles' as one item; it is two suggestions",
                              to=["(omit)", "handles"]),
    "minimize": fix("set", "the source's second suggestion, 'method', looks like a mistake and is left out",
                    to=["decrease"]),
    "notwithstanding": fix("set", "the source's 'inspite of' is spelled 'in spite of'", to=["in spite of", "still"]),
    "time period": fix("set", "the source's '(either one)' means 'time' or 'period'", to=["time", "period"]),
    "this activity, command": fix("set", "the source means 'this activity' or 'this command'",
                                  froms=["this activity", "this command"]),
}


def parse_gsa(text: str, used: set) -> list[tuple[list[str], list[str], str | None]]:
    out = []
    in_table = False
    rows = 0
    for line in text.splitlines():
        if not in_table:
            if re.match(r"^\*\*Don.t say\*\*\s*\|\s*\*\*Say\*\*\s*$", line.strip()):
                in_table = True
            continue
        if line.startswith("{:"):
            break
        if re.fullmatch(r"-+\s*\|\s*-+", line.strip()):
            continue
        if "|" not in line:
            raise BuildError(f"gsa: unexpected line inside the table: {line!r}")
        rows += 1
        left, right = line.split("|", 1)
        left = clean(left.replace("**", "").replace("\\_", "_"))
        right = clean(right.replace("**", ""))
        froms = [f.strip() for f in left.split(", ")]
        to = [t.strip() for t in right.split(", ")]
        if left in GSA_FIXES:
            used.add(left)
            f = GSA_FIXES[left]
            if f["action"] == "drop":
                continue
            froms = f["froms"] or froms
            to = f["to"] or to
        out.append((froms, to, None))
    if rows == 0:
        raise BuildError("gsa: no table found")
    return out


def parse_18f(text: str) -> list[tuple[list[str], list[str], str | None]]:
    m = re.search(r"^## Words to avoid\s*$(.*?)^## ", text, re.M | re.S)
    if not m:
        raise BuildError("18f: no 'Words to avoid' section")
    items = [line[2:].strip() for line in m.group(1).splitlines() if line.startswith("- ")]
    if not items:
        raise BuildError("18f: the 'Words to avoid' section has no list")
    out = []
    for item in items:
        im = re.fullmatch(r"(.*?)(?:\s*\((.*)\))?", item)
        head, note = im.group(1), im.group(2) or ""
        froms = [clean(f) for f in head.split(" or ")]
        italics = [clean(x) for x in re.findall(r"_([^_]+)_", note)]
        plain = clean(note.replace("_", ""))
        unless, to = None, []
        omit = re.compile(r"\bomit\b|\bleave\b.*\bout\b")
        if plain.startswith("unless "):
            rest = plain[len("unless "):]
            if ", use " in rest:
                unless, use = rest.split(", use ", 1)
                to = [x for x in italics if x in use] + (["(omit)"] if omit.search(use) else [])
            else:
                unless = rest
        elif plain.startswith("use "):
            to = italics + (["(omit)"] if omit.search(plain) else [])
        elif plain:
            to = [f"({plain})"]
        out.append((froms, to, unless))
    return out


GOVUK_FIXES = {
    "commit/pledge, use 'plan to x', or 'we're going to x' where 'x' is a specific verb":
        fix("set", "the source uses 'x' for a specific verb; written here as '(verb)'",
            to=["plan to (verb)", "we're going to (verb)"]),
    "facilitate, say something specific about how you're helping - for example, use 'run' if talking about a workshop":
        fix("advice", "'run' is the source's example for one case (a workshop), not a general swap, so the "
                      "whole advice is kept instead"),
    "going/moving forward, use 'from now on' or 'in the future' (it's unlikely we are giving travel directions)":
        fix("set", "the source means 'going forward' and 'moving forward'",
            froms=["going forward", "moving forward"]),
    "hub, portal or one-stop shop, use 'website' or 'service' (unless those words are in the name of the website or service)":
        fix("set", "three words share one line, and the 'unless' note comes at the end",
            froms=["hub", "portal", "one-stop shop"], to=["website", "service"],
            unless="those words are in the name of the website or service"),
}


class _GovukItems(html.parser.HTMLParser):
    """The text of each <li> after <h3 id="words-to-avoid">, up to the next <h3>."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.active = False
        self.done = False
        self.in_li = False
        self.items: list[str] = []
        self.buf: list[str] = []

    def handle_starttag(self, tag, attrs):
        if self.done:
            return
        if tag == "h3":
            if self.active:
                self.done = True
            elif dict(attrs).get("id") == "words-to-avoid":
                self.active = True
        elif tag == "li" and self.active:
            self.in_li, self.buf = True, []

    def handle_endtag(self, tag):
        if tag == "li" and self.in_li:
            self.items.append("".join(self.buf))
            self.in_li = False

    def handle_data(self, data):
        if self.in_li:
            self.buf.append(data)


QUOTED = re.compile(r"‘(.+?)’(?![^\W\d_])")


def parse_govuk(text: str, used: set) -> list[tuple[list[str], list[str], str | None]]:
    p = _GovukItems()
    p.feed(text)
    if not p.items:
        raise BuildError("govuk: no 'Words to avoid' list")
    out = []
    for raw in p.items:
        quoted = [clean(q) for q in QUOTED.findall(raw) if len(clean(q)) > 1]
        item = clean(raw)
        parts = split_top(item, ", ")
        head, rest = parts[0], ", ".join(parts[1:])
        hm = re.fullmatch(r"(.*?)(?:\s*\(unless (.*)\))?", head)
        name, unless = hm.group(1), hm.group(2)
        froms = [clean(f) for f in re.split(r"/| or ", name)]
        omit = bool(re.search(r"usually not needed|usually superfluous", rest))
        to = (["(omit)"] if omit else []) + quoted
        if not to:
            to = [f"({rest})"]
        if item in GOVUK_FIXES:
            used.add(item)
            f = GOVUK_FIXES[item]
            if f["action"] == "advice":
                to = [f"({rest})"]
            else:
                froms = f["froms"] or froms
                to = f["to"] or to
                unless = f["unless"] or unless
        out.append((froms, to, unless))
    return out


def _block(text: str, start: str, end: str, what: str) -> str:
    m = re.search(start + r"(.*?)" + end, text, re.S)
    if not m:
        raise BuildError(f"iso-24495: {what} not found")
    return m.group(1)


def parse_iso(lexicon: str, audit: str, skipped: list) -> list[tuple[list[str], list[str], str | None]]:
    out = []
    block = _block(lexicon, r"export const COMPLEX_WORDS[^=]*=\s*new Map\(\[", r"\]\);", "COMPLEX_WORDS")
    pairs = re.findall(r'\[\s*"([^"]+)"\s*,\s*"([^"]+)"\s*\]', block)
    if len(pairs) != block.count("["):
        raise BuildError("iso-24495: COMPLEX_WORDS has an entry this script cannot read")
    out += [([a], [b], None) for a, b in pairs]

    block = _block(audit, r"const WORDY_PHRASES[^=]*=\s*\[", r"\]\.sort", "WORDY_PHRASES")
    pairs = re.findall(r'\{\s*phrase:\s*"([^"]+)"\s*,\s*lean:\s*"([^"]+)"\s*\}', block)
    if len(pairs) != block.count("{"):
        raise BuildError("iso-24495: WORDY_PHRASES has an entry this script cannot read")
    out += [([a], [b], None) for a, b in pairs]

    block = _block(audit, r"const DOUBLETS[^=]*=\s*\[", r"\]\.sort", "DOUBLETS")
    entries = re.findall(r'\{\s*phrase:\s*"([^"]+)"\s*(?:,\s*lean:\s*"([^"]+)"\s*)?(?:,\s*legalRegister:\s*true\s*)?\}', block)
    if len(entries) != block.count("{"):
        raise BuildError("iso-24495: DOUBLETS has an entry this script cannot read")
    for phrase, lean in entries:
        if not lean:
            skipped.append((SRC_ISO, phrase, "a legal phrase with no plainer replacement in the source "
                                              "(it only asks the writer to check the phrase suits the reader)"))
            continue
        out.append(([phrase], [lean], None))
    return out


def merge(lists: list[tuple[str, list]]) -> list[dict]:
    merged: dict[str, dict] = {}
    for source, rows in lists:
        for froms, to, unless in rows:
            for f in froms:
                key = clean(f).lower()
                if not key:
                    raise BuildError(f"{source}: an empty word")
                e = merged.setdefault(key, {"to": [], "sources": [], "unless": []})
                for t in to:
                    t = clean(t)
                    if t and t.lower() not in (x.lower() for x in e["to"]):
                        e["to"].append(t)
                if source not in e["sources"]:
                    e["sources"].append(source)
                if unless:
                    u = clean(unless)
                    if u not in e["unless"]:
                        e["unless"].append(u)
    out = []
    for key in sorted(merged):
        e = merged[key]
        entry = {"from": key, "to": e["to"], "source": ", ".join(e["sources"])}
        if e["unless"]:
            entry["unless"] = "; ".join(e["unless"])
        out.append(entry)
    return out


# ---------------------------------------------------------------------------------------------
# Technical-exceptions list: jargone's own. A word here is one whose plain swap would change a
# precise meaning in software, law, medicine, statistics or finance. J07 does not suggest a swap for
# it; J01 still asks for it to be defined if the reader may not know it. Google's developer
# documentation style guide word list (CC BY 4.0) was used only as a reference for which words are
# risky; none of its text is here. Base forms, plus the exact form a replacement entry uses where it
# differs (for example "terminates", "parameters", "impacted").
# ---------------------------------------------------------------------------------------------

EXCEPTIONS = {
    # Software
    "address": "software: a memory, network or email address",
    "allocate": "software: set aside memory or another resource",
    "allocation": "software: memory or another resource set aside",
    "argument": "software: a value passed to a function; law: a case put to a court",
    "capability": "software: a permission the system grants (Linux capabilities)",
    "commit": "software: record a change in version control or a database",
    "component": "software: a reusable part of an interface or system",
    "counter": "software: a variable that counts",
    "delete": "software: remove data (in SQL, DELETE and DROP do different things)",
    "deliver": "software: message delivery has exact guarantees (at least once, at most once)",
    "deploy": "software: put a release into an environment",
    "deployment": "software: a release put into an environment",
    "deprecate": "software: mark as due to be removed",
    "deprecated": "software: marked as due to be removed",
    "drive": "software: a storage device",
    "driver": "software: the program that runs a device",
    "enable": "software: turn on a feature or setting",
    "executable": "software: a file the computer can run",
    "execute": "software: run code; law: sign a document so it takes effect",
    "execution": "software: running code; law: signing a document",
    "forward": "software: pass traffic on (port forwarding)",
    "function": "software: a named piece of code; maths: a mapping from inputs to outputs",
    "hub": "software: a network device",
    "implement": "software: write the code that fulfils an interface or design",
    "implementation": "software: the code that fulfils an interface or design",
    "initial": "software: the starting value or state (initial state)",
    "initialize": "software: set the starting value",
    "initialise": "software: set the starting value",
    "instance": "software: one object, process or server made from a template",
    "interface": "software: where two parts of a system meet (API, user interface)",
    "invoke": "software: call a function or a command",
    "key": "software: a value that unlocks, signs or looks up data (API key, primary key)",
    "maintain": "software: keep a project working and up to date",
    "maintainer": "software: the person responsible for a project",
    "maintenance": "software: work to keep a system working",
    "monitor": "software: collect measurements from a running system",
    "monitoring": "software: collecting measurements from a running system",
    "option": "software: a setting passed to a command; finance: a contract to buy or sell",
    "parameter": "software: a named input to a function, model or command",
    "parameters": "software: named inputs to a function, model or command",
    "persist": "software: store data so it survives a restart",
    "render": "software: draw a page, an image or a template",
    "request": "software: a message sent to a server (an HTTP request)",
    "require": "software: load a module; specifications: a 'must' (RFC 2119)",
    "required": "software: a field or setting that must be given",
    "requirement": "software: something a system must do",
    "requirements": "software: what a system must do",
    "terminate": "software: end a process (a stop signal only pauses it); law: end a contract",
    "terminates": "software: ends a process; law: ends a contract",
    "termination": "law: the ending of a contract",
    "transform": "maths and data work: a defined change (Fourier transform, data transform)",
    "transforming": "maths and data work: applying a defined change",
    "type": "software: a data type",
    "utilization": "software and operations: the share of capacity in use (CPU utilization)",
    "validate": "software: check data against rules",
    "validation": "software: checking data against rules",
    # Law
    "appropriate": "law and public finance: as a verb, set money aside by law",
    "breach": "law: failure to keep a contract or a duty",
    "consideration": "law: what each party gives to make a contract binding",
    "damages": "law: money a court awards",
    "deem": "law: treat as true by law, whatever the facts",
    "deemed": "law: treated as true by law, whatever the facts",
    "disclose": "law, finance and security: a duty to reveal information",
    "disclosure": "law, finance and security: revealing information as required",
    "employ": "law: hire someone (employment law)",
    "equitable": "law: decided under equity, a branch of law",
    "exhibit": "law: a document attached to a contract or filed in a case",
    "forfeit": "law: lose a right or property as a penalty",
    "forfeiture": "law: losing a right or property as a penalty",
    "indemnify": "law: promise to cover another's loss",
    "indemnity": "law: a promise to cover another's loss",
    "jurisdiction": "law: a court's power over a case or a place",
    "liability": "law and finance: a legal duty or a debt",
    "material": "law and finance: important enough to change a decision (a material breach)",
    "notice": "law: a formal message with legal effect",
    "party": "law: a person or body bound by a contract or in a case",
    "permit": "law: a licence to do something",
    "practicable": "law: possible to do (not the same as practical)",
    "remedy": "law: what a court orders to put a wrong right",
    "shall": "law and standards: states a duty",
    "void": "law: having no legal effect",
    "waive": "law: give up a right on purpose",
    "waiver": "law: giving up a right on purpose",
    "warrant": "law: a court order; finance: a right to buy shares",
    "warranty": "law: a promise in a contract",
    "witnessed": "law: signed in front of a witness",
    # Medicine
    "acute": "medicine: sudden and short (not the same as severe)",
    "administer": "medicine: give a drug or treatment",
    "benign": "medicine: not cancer",
    "chronic": "medicine: lasting a long time",
    "complication": "medicine: a new problem caused by an illness or treatment",
    "dose": "medicine: the amount of a drug taken at one time",
    "impacted": "medicine: stuck in place (an impacted tooth)",
    "indication": "medicine: a reason to use a treatment",
    "intervention": "medicine and research: a treatment or action being tested",
    "malignant": "medicine: cancer that can spread",
    "negative": "medicine: a test result that did not find what it looked for",
    "positive": "medicine: a test result that found what it looked for",
    "presentation": "medicine: how an illness shows itself in a patient",
    "subject": "medicine and research: a person in a study; software: the subject of a certificate or message",
    "viable": "medicine and biology: able to live or grow",
    # Statistics and maths
    "accuracy": "statistics: closeness to the true value (not the same as precision)",
    "accurate": "statistics: close to the true value (not the same as precise)",
    "bias": "statistics: a systematic error",
    "confidence": "statistics: as in a confidence interval",
    "correlation": "statistics: how strongly two measures move together",
    "error": "statistics: the gap between a measured and a true value",
    "estimate": "statistics: a value worked out from a sample",
    "feasible": "maths: meeting every constraint (a feasible solution)",
    "likelihood": "statistics: how well a model explains the data seen",
    "magnitude": "maths and physics: the size of a quantity; an order of magnitude",
    "maximum": "maths: the largest value a function takes",
    "mean": "statistics: the arithmetic average",
    "median": "statistics: the middle value",
    "minimize": "maths: find the smallest value (optimisation)",
    "minimum": "maths: the smallest value a function takes",
    "normal": "statistics: following the normal (bell-shaped) distribution",
    "objective": "maths: the quantity being optimised (an objective function)",
    "observation": "statistics: one recorded data value",
    "observe": "statistics: record a data value; software: watch a system's behaviour (observability)",
    "observed": "statistics: recorded in the data",
    "odds": "statistics: the chance of an event divided by the chance it does not happen",
    "optimum": "maths: the best value under stated constraints",
    "population": "statistics: the whole group a sample is drawn from",
    "power": "statistics: the chance a test detects a real effect",
    "precise": "statistics: giving close results when repeated (not the same as accurate)",
    "precision": "statistics: how close repeated results are",
    "probability": "statistics: a measure of chance between 0 and 1",
    "random": "statistics: chosen by chance, each with a known probability",
    "regression": "statistics: a model of how one measure depends on others",
    "reliability": "statistics: how consistent a measurement is",
    "remainder": "maths: what is left after division",
    "robust": "statistics: little affected by outliers or broken assumptions",
    "sample": "statistics: the part of a population that was measured",
    "selection": "statistics: as in selection bias",
    "significance": "statistics: how unlikely a result is under chance alone",
    "significant": "statistics: unlikely to be due to chance alone (a test result)",
    "sufficient": "maths and statistics: a sufficient condition; a sufficient statistic",
    "valid": "logic and measurement: follows from its premises; measures what it claims to",
    "validity": "measurement: whether a test measures what it claims to",
    "variance": "statistics: the spread of values around the mean",
    # Finance
    "accrual": "finance: income or cost recorded before money changes hands",
    "accrue": "finance: build up over time before it is paid (accrued interest)",
    "asset": "finance: something of value that is owned",
    "benefit": "finance and insurance: a payment under a policy or scheme (a death benefit)",
    "bond": "finance: a loan sold as a security",
    "capital": "finance: money used to run or grow a business",
    "consolidate": "finance: combine group accounts into one set (consolidated accounts)",
    "consolidated": "finance: combined into one set of group accounts",
    "credit": "finance: money lent, or an entry that adds to an account",
    "dividend": "finance: a share of profits paid to owners",
    "equity": "finance: ownership value; law: a branch of law",
    "exposure": "finance: the amount that could be lost",
    "hedge": "finance: a deal that offsets a risk",
    "interest": "finance: money paid for borrowing; law: a right in property",
    "leverage": "finance: the use of borrowed money",
    "liquidity": "finance: how easily something turns into cash",
    "margin": "finance: profit as a share of sales, or money held against a trade",
    "premium": "finance and insurance: the price of a policy or an option",
    "principal": "finance: the amount borrowed, before interest",
    "provision": "finance: money set aside for a likely cost; law: a clause",
    "return": "finance: the gain or loss on an investment",
    "security": "finance: a tradable asset such as a share or a bond",
    "yield": "finance: income from an investment as a share of its price",
}


def check_exceptions() -> list[str]:
    bad = [w for w in EXCEPTIONS if not re.fullmatch(r"[a-z]+", w)]
    if bad:
        raise BuildError(f"exceptions must be single lower-case words: {bad}")
    return sorted(EXCEPTIONS)


# ---------------------------------------------------------------------------------------------
# Output
# ---------------------------------------------------------------------------------------------


def template_literal(s: str) -> str:
    """The body of a JavaScript template literal whose value is s. Non-ASCII becomes \\u escapes."""
    s = s.replace("\\", "\\\\").replace("`", "\\`").replace("${", "\\${")
    return "".join(c if ord(c) < 128 else f"\\u{{{ord(c):x}}}" for c in s)


def render_data_ts(date: str, fam: list[tuple[str, str]], reps: list[dict], exc: list[str]) -> str:
    fam_body = "\n".join(template_literal(w) + "\\t" + r for w, r in fam)
    rep_body = template_literal("[\n" + ",\n".join(json.dumps(e, ensure_ascii=False) for e in reps) + "\n]")
    exc_body = template_literal("\n".join(exc))
    return (
        f"// GENERATED by scripts/build_data.py on {date}. Do not edit by hand: change the script and run it again.\n"
        "// Where each list comes from, its licence and the credit it requires: data/NOTICE.md. The plugin's code\n"
        "// is MIT; these lists keep their own licences.\n"
        "//\n"
        f"// FAMILIARITY: {len(fam):,} lines \"word<TAB>rating\". One lower-case word per line (letters, and\n"
        "//   apostrophes inside a word); rating from 1 (unknown) to 7 (very familiar), one decimal place.\n"
        "//   Forms ending in 's are not listed: remove a final 's before looking a word up. A word that is\n"
        "//   not listed was rated lowest by the source's authors (about 3.5 or below) or not rated at all:\n"
        "//   treat it as unfamiliar.\n"
        f"// REPLACEMENTS: a JSON array of {len(reps)} entries {{ from, to, source, unless? }}. `from` is lower\n"
        "//   case and may be a phrase. A `to` item in round brackets is advice, not a word to paste in;\n"
        "//   \"(omit)\" means leave it out. `to` may be empty when the source gives no replacement.\n"
        "//   `source` names every source that lists it, joined with \", \".\n"
        f"// EXCEPTIONS: {len(exc)} lines, one lower-case word each: jargone's own technical-exceptions list.\n"
        "\n"
        f"export const FAMILIARITY: string = `{fam_body}`\n"
        "\n"
        f"export const REPLACEMENTS: string = `{rep_body}`\n"
        "\n"
        f"export const EXCEPTIONS: string = `{exc_body}`\n"
    )


def render_notice_block(date: str, got: dict, fam_stats: dict, per_source: dict, reps: list[dict],
                        exc: list[str], changes: list[tuple[str, str, str]], skipped: list) -> str:
    L = []
    L.append(f"Built on {date} by `scripts/build_data.py`. This part of the file is written by the script; "
             "do not edit it by hand.")
    L.append("")
    L.append("| Source | Address | Got | Fingerprint (SHA-256) |")
    L.append("|---|---|---|---|")
    for key, src in SOURCES.items():
        g = got[key]
        L.append(f"| {src['name']} | <{src['url']}> | {g['how']} | `{g['sha256']}` |")
    L.append("")
    s = fam_stats
    L.append(f"**Familiarity list:** {s['rows']:,} rows read; {s['kept']:,} words kept, rated {s['lowest']} "
             f"to {s['highest']}; {s['at_or_above_6']:,} of them at 6.0 or above (the default cut-off). "
             f"Left out: {s['phrases']:,} phrases, {s['hyphenated']:,} hyphenated words, {s['not_letters']:,} "
             f"entries not made of letters (numbers, abbreviations with dots, slashes, underscores), "
             f"{s['ends_in_s']:,} forms ending in 's, and {s['other_type']:,} rows of another type. "
             f"{s['same_word_other_case']:,} words appeared twice once lower-cased; one of each was kept.")
    L.append("")
    counts = "; ".join(f"{name} {n}" for name, n in per_source.items())
    L.append(f"**Replacement list:** {len(reps)} entries after merging. Entries read from each source "
             f"(a line naming several words counts once per word): {counts}. "
             f"{sum(1 for e in reps if ', ' in e['source'])} entries are listed by more than one source.")
    L.append("")
    L.append(f"**Technical-exceptions list:** {len(exc)} words.")
    overlap = [e["from"] for e in reps if e["from"] in set(exc)]
    L.append("")
    L.append(f"**Replacement-list words made exceptions** ({len(overlap)}), because their plain swap would "
             "damage technical writing (the reason for each is next to the word in `scripts/build_data.py`): "
             + ", ".join(overlap) + ".")
    L.append("")
    L.append("**Replacement-list entries this script changed or left out:**")
    L.append("")
    for src, line, reason in changes:
        L.append(f"- {src}, \"{line}\": {reason}.")
    for src, line, reason in skipped:
        L.append(f"- {src}, \"{line}\": left out; {reason}.")
    return "\n".join(L)


def fetch(key: str, path: str | None, today: str) -> tuple[bytes, dict]:
    src = SOURCES[key]
    if path:
        data = Path(path).read_bytes()
        how = f"local copy ({src['flag']})"
    else:
        req = urllib.request.Request(src["url"], headers={"User-Agent": "jargone-build-data/1"})
        with urllib.request.urlopen(req, timeout=120) as r:
            data = r.read()
        how = f"downloaded {today}"
    return data, {"how": how, "sha256": sha256(data), "bytes": len(data)}


def main(argv: list[str] | None = None) -> int:
    here = Path(__file__).resolve().parent
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    for key, src in SOURCES.items():
        ap.add_argument(src["flag"], dest=key.replace("-", "_"), metavar="PATH",
                        help=f"use this local file instead of downloading {src['url']}")
    ap.add_argument("--date", help="build date to record (YYYY-MM-DD); default today")
    ap.add_argument("--root", default=str(here.parent), help="the plugin folder to write into (default: %(default)s)")
    args = ap.parse_args(argv)

    today = args.date or _dt.date.today().isoformat()
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", today):
        raise BuildError("--date must be YYYY-MM-DD")
    root = Path(args.root)
    notice_path = root / "data" / "NOTICE.md"
    notice = notice_path.read_text(encoding="utf-8")
    if "<!-- build:start -->" not in notice or "<!-- build:end -->" not in notice:
        raise BuildError(f"{notice_path} needs the <!-- build:start --> and <!-- build:end --> markers")

    raw, got = {}, {}
    for key in SOURCES:
        raw[key], got[key] = fetch(key, getattr(args, key.replace("-", "_")), today)
        print(f"{key}: {got[key]['how']}, {got[key]['bytes']} bytes, sha256 {got[key]['sha256']}")

    fam, fam_stats = build_familiarity(raw["familiarity"])

    text = {k: raw[k].decode("utf-8") for k in SOURCES if k != "familiarity"}
    used: set = set()
    skipped: list = []
    gsa = parse_gsa(text["gsa"], used)
    eighteen_f = parse_18f(text["18f"])
    govuk = parse_govuk(text["govuk"], used)
    iso = parse_iso(text["iso-lexicon"], text["iso-audit"], skipped)
    unused = [k for k in list(GSA_FIXES) + list(GOVUK_FIXES) if k not in used]
    if unused:
        raise BuildError("these fixes no longer match a source line, so the source changed; read it again "
                         "and update the fixes: " + "; ".join(repr(u) for u in unused))
    changes = [(SRC_GSA, k, f["reason"]) for k, f in GSA_FIXES.items()] + \
              [(SRC_GOVUK, k, f["reason"]) for k, f in GOVUK_FIXES.items()]

    lists = [(SRC_GSA, gsa), (SRC_18F, eighteen_f), (SRC_GOVUK, govuk), (SRC_ISO, iso)]
    per_source = {name: sum(len(froms) for froms, _, _ in rows) for name, rows in lists}
    for name, n in per_source.items():
        if n == 0:
            raise BuildError(f"{name}: no entries read")
    reps = merge(lists)
    exc = check_exceptions()

    out_ts = root / "hooks" / "data.ts"
    out_ts.parent.mkdir(parents=True, exist_ok=True)
    ts = render_data_ts(today, fam, reps, exc)
    out_ts.write_text(ts, encoding="utf-8")

    block = render_notice_block(today, got, fam_stats, per_source, reps, exc, changes, skipped)
    notice = re.sub(r"<!-- build:start -->.*?<!-- build:end -->",
                    lambda _: "<!-- build:start -->\n" + block + "\n<!-- build:end -->", notice, flags=re.S)
    notice_path.write_text(notice, encoding="utf-8")

    print(f"familiarity: {fam_stats['rows']} rows read, {fam_stats['kept']} words kept "
          f"({fam_stats['at_or_above_6']} at 6.0 or above)")
    print(f"replacements: {len(reps)} entries; read per source: {per_source}")
    print(f"exceptions: {len(exc)} words; {sum(1 for e in reps if e['from'] in set(exc))} are also replacement-list entries")
    print(f"wrote {out_ts} ({len(ts.encode('utf-8'))} bytes) and the build part of {notice_path}")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except BuildError as e:
        print(f"build_data: {e}", file=sys.stderr)
        sys.exit(1)
